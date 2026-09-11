"""v3: 按附件 5 模板重写 result2.xlsx / result4-2.xlsx（修正后的信息结构），并导出表 1/2/3 格式表格。"""
from __future__ import annotations

import os
import sys

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import model as M
import pipeline as P
from write_results import slot_labels, block_rows

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "out")
DEST = os.path.join(os.path.dirname(ROOT), "outputs", "result_files")
os.makedirs(DEST, exist_ok=True)
DATES4 = ["2025-03-20", "2025-06-21", "2025-09-23", "2025-12-21"]
WIN = [60, 72, 84, 96, 108, 120]
BLK = [("0:00-4:00", 0, 24), ("4:00-8:00", 24, 48), ("8:00-12:00", 48, 72),
       ("12:00-16:00", 72, 96), ("16:00-20:00", 96, 120),
       ("20:00-24:00", 120, 144)]


def segs(e):
    out, k = [], 0
    while k < 144:
        if e[k] > 1e-6:
            k0 = k
            while k < 144 and e[k] > 1e-6:
                k += 1
            out.append((k0, k - 1, float(e[k0:k].sum())))
        else:
            k += 1
    return out


def main():
    import openpyxl
    data = P.load_data()
    dates = [str(x) for x in data["dates2"]]
    idx = P._day_range(data)
    labels = slot_labels()
    blocks = block_rows()
    p_fix = data["att1"][:, 0]
    T = []

    for fname, npz, price_mode, tag in [
            ("result2.xlsx", "p2_v3.npz", "att1", "问题 2"),
            ("result4-2.xlsx", "p42_v3.npz", "att4", "问题 4-2")]:
        d = np.load(os.path.join(OUT, npz), allow_pickle=True)
        dd = [str(x) for x in d["dates"]]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "计划购电量"
        ws.append(["日期\\时间"] + labels + ["全天购电量", "全天购电费"])
        for j, dt in enumerate(dd):
            i = dates.index(dt)
            p = p_fix if price_mode == "att1" else data["price_rt"][i]
            b = d["plan_b"][j]
            ws.append([dt] + [round(float(v), 4) for v in b]
                      + [round(float(b.sum()), 4),
                         round(float(np.sum(p * b)), 4)])
        ws = wb.create_sheet("充放电量")
        ws.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
        for j, dt in enumerate(dd):
            for bi, (name, a, b) in enumerate(blocks):
                ws.append([dt if bi == 0 else None, name,
                           round(float(d["c"][j][a:b].sum()), 4),
                           round(float(d["d"][j][a:b].sum()), 4),
                           "0:00" if bi == 0 else ("24:00" if bi == 1 else None),
                           6000.0 if bi == 0 else (round(float(d["s"][j][-1]), 4)
                                                   if bi == 1 else None)])
        ws = wb.create_sheet("紧急购电量")
        ws.append(["日期", "购电时间段", "购电量"])
        for j, dt in enumerate(dd):
            e = d["emergency"][j]
            S = segs(e)
            if not S:
                ws.append([dt, "无", 0.0])
            for t, (a, b, v) in enumerate(S):
                ws.append([dt if t == 0 else None,
                           f"{labels[a].split('-')[0]}—{labels[b].split('-')[1]}",
                           round(v, 4)])
        wb.save(os.path.join(DEST, fname))
        print(f"写出 {fname}（{len(dd)} 天，紧急购电段数合计 "
              f"{sum(len(segs(d['emergency'][j])) for j in range(len(dd)))}）")

        # 表 1 / 表 2 / 表 3 格式（4 个指定日期）
        T.append(f"\n**{tag}：表 1 格式（指定时间段购电量，kWh）**\n")
        T.append("| 日期 | " + " | ".join(labels[k] for k in WIN) +
                 " | 全天购电量 | 全天购电费(元) |")
        T.append("|" + "---|" * (len(WIN) + 3))
        for dt in DATES4:
            j = dd.index(dt)
            i = dates.index(dt)
            p = p_fix if price_mode == "att1" else data["price_rt"][i]
            b = d["plan_b"][j]
            T.append("| %s | %s | %.2f | %.2f |" % (
                dt, " | ".join("%.2f" % b[k] for k in WIN), b.sum(),
                float(np.sum(p * b))))
        T.append(f"\n**{tag}：表 2 格式（储能充放电量与 0:00/24:00 储电量，kWh）**\n")
        T.append("| 日期 | 时段 | 充电量 | 放电量 | 时段 | 充电量 | 放电量 | 0:00 储电量 | 24:00 储电量 |")
        T.append("|---|---|---|---|---|---|---|---|---|")
        for dt in DATES4:
            j = dd.index(dt)
            c, dd_ = d["c"][j], d["d"][j]
            for t, ((n1, a1, b1), (n2, a2, b2)) in enumerate(zip(BLK[:3], BLK[3:])):
                T.append("| %s | %s | %.2f | %.2f | %s | %.2f | %.2f | %s | %s |" % (
                    dt if t == 0 else "", n1, c[a1:b1].sum(), dd_[a1:b1].sum(),
                    n2, c[a2:b2].sum(), dd_[a2:b2].sum(),
                    "6000.00" if t == 0 else "",
                    "%.2f" % d["s"][j][-1] if t == 0 else ""))
        T.append(f"\n**{tag}：表 3 格式（紧急购电，kWh）**\n")
        T.append("| 日期 | 紧急购电时段与电量 | 合计 | 紧急购电费(元) |")
        T.append("|---|---|---|---|")
        for dt in DATES4:
            j = dd.index(dt)
            e = d["emergency"][j]
            i = dates.index(dt)
            p = p_fix if price_mode == "att1" else data["price_rt"][i]
            S = segs(e)
            txt = "；".join("%s—%s: %.1f" % (labels[a].split("-")[0],
                                             labels[b].split("-")[1], v)
                            for a, b, v in S) or "无"
            T.append("| %s | %s | %.1f | %.0f |" % (
                dt, txt, e.sum(), float(np.sum(5.0 * p * e))))
    txt = "\n".join(T)
    open(os.path.join(OUT, "v3_tables.md"), "w", encoding="utf-8").write(txt)
    print(txt)


if __name__ == "__main__":
    main()
