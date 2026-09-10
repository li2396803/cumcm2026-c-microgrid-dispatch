"""写出 v2 结果文件 (按附件 5 模板): result2_v2 / result3_v2 / result4-3_v2."""
from __future__ import annotations

import os
import sys

import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import model as M
import pipeline as P
from write_results import slot_labels, block_rows

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
OUT = os.path.join(ROOT, "out")
DEST = os.path.join(os.path.dirname(ROOT), "outputs", "result_files")
os.makedirs(DEST, exist_ok=True)


def main():
    import openpyxl
    data = P.load_data()
    dates = [str(x) for x in data["dates2"]]
    idx = P._day_range(data)
    labels = slot_labels()
    blocks = block_rows()
    p_fix = data["att1"][:, 0]
    prt = data["price_rt"]

    # ---------------- result2_v2: 寿命感知 (c_deg=0.104) 的完全信息策略 ----------------
    B = np.zeros((len(idx), 144))
    C = np.zeros((len(idx), 144))
    D = np.zeros((len(idx), 144))
    S = np.zeros((len(idx), 144))
    for j, i in enumerate(idx):
        sol = M.solve_day(p_fix, data["load"][i] * M.DT, data["pv"][i] * M.DT,
                          deg_cost=0.104, solver="highs-ds")
        B[j], C[j], D[j], S[j] = sol["b"], sol["c"], sol["d"], sol["s"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "计划购电量"
    ws.append(["日期\\时间"] + labels + ["全天购电量", "全天购电费"])
    for j, i in enumerate(idx):
        cost = float(np.sum(p_fix * B[j]))
        ws.append([dates[i]] + [round(float(v), 4) for v in B[j]]
                  + [round(float(B[j].sum()), 4), round(cost, 4)])
    ws = wb.create_sheet("充放电量")
    ws.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
    for j, i in enumerate(idx):
        for bi, (name, a, b) in enumerate(blocks):
            ws.append([dates[i] if bi == 0 else None, name,
                       round(float(C[j][a:b].sum()), 4),
                       round(float(D[j][a:b].sum()), 4),
                       "0:00" if bi == 0 else ("24:00" if bi == 1 else None),
                       6000.0 if bi == 0 else (round(float(S[j][-1]), 4)
                                               if bi == 1 else None)])
    ws = wb.create_sheet("紧急购电量")
    ws.append(["日期", "购电时间段", "购电量"])
    for j, i in enumerate(idx):
        ws.append([dates[i], "无", 0.0])
    wb.save(os.path.join(DEST, "result2_v2.xlsx"))

    # ---------------- result3_v2 / result4-3_v2: MPC 闭环 ----------------
    def write_mpc(fname, mpc_file, price_mode):
        m = np.load(os.path.join(OUT, mpc_file), allow_pickle=True)
        dd = list(m["dates"])
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "计划购电量"
        ws.append(["日期\\时间"] + labels + ["全天购电量", "全天购电费"])
        for j, d in enumerate(dd):
            i = dates.index(str(d))
            p = p_fix if price_mode == "att1" else prt[i]
            bp = m["plan"][j]
            ws.append([str(d)] + [round(float(v), 4) for v in bp]
                      + [round(float(bp.sum()), 4),
                         round(float(np.sum(p * bp)), 4)])
        ws = wb.create_sheet("调整购电量")
        ws.append(["日期\\时间"] + labels + ["全天购电量", "全天购电费"])
        for j, d in enumerate(dd):
            b = m["b"][j]
            ws.append([str(d)] + [round(float(v), 4) for v in b]
                      + [round(float(b.sum()), 4),
                         round(float(m["cost"][j]), 4)])
        ws = wb.create_sheet("充放电量")
        ws.append(["日期", "时间段", "充电量", "放电量", "时刻", "储电量"])
        for j, d in enumerate(dd):
            for bi, (name, a, b) in enumerate(blocks):
                ws.append([str(d) if bi == 0 else None, name,
                           round(float(m["c"][j][a:b].sum()), 4),
                           round(float(m["d"][j][a:b].sum()), 4),
                           "0:00" if bi == 0 else ("24:00" if bi == 1 else None),
                           6000.0 if bi == 0 else None])
        ws = wb.create_sheet("紧急购电量")
        ws.append(["日期", "购电时间段", "购电量"])
        for j, d in enumerate(dd):
            e = m["emergency"][j]
            k = 0
            rows = []
            while k < 144:
                if e[k] > 1e-6:
                    k0 = k
                    while k < 144 and e[k] > 1e-6:
                        k += 1
                    rows.append((k0, k - 1, float(e[k0:k].sum())))
                else:
                    k += 1
            if not rows:
                ws.append([str(d), "无", 0.0])
            for t, (a, b, v) in enumerate(rows):
                ws.append([str(d) if t == 0 else None,
                           f"{labels[a].split('-')[0]}—{labels[b].split('-')[1]}",
                           round(v, 4)])
        wb.save(os.path.join(DEST, fname))
        print("wrote", fname, "天数", len(dd))

    write_mpc("result3_v2.xlsx", "v2_mpc_b.npz", "att1")
    if os.path.exists(os.path.join(OUT, "v2_mpc_att4.npz")):
        write_mpc("result4-3_v2.xlsx", "v2_mpc_att4.npz", "att4")
    print("done ->", DEST)


if __name__ == "__main__":
    main()
